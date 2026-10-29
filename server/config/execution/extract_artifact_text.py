"""Extracts text from binary artifacts for dsh-work (execution isolation plan 4.6, 2b-3).

Runs in a disposable container from the execution image under the Python sandbox posture: no
network, read-only root, memory, CPU and time limits. Files produced in an execution container
may be crafted to attack a parser, so they are never parsed in the service process.

Usage: python -I -B extract_artifact_text.py <name>...
Reads /input/<name> and writes /output/<index>.json (metadata) and, when text was found,
/output/<index>.txt (UTF-8). The service validates and caps everything it reads back.
"""

import datetime
import itertools
import json
import os
import re
import subprocess
import sys

MAX_CHARACTERS = 512_000
MAX_SHEETS = 30
MAX_CELLS = 20_000
MAX_SPREADSHEET_ROWS = 100_000
MAX_PDF_PAGES = 200
MAX_SLIDES = 300
TOOL_TIMEOUT_SECONDS = 60


class Unusable(Exception):
    """The file parsed but cannot become reference text; `code` is reported to the service."""

    def __init__(self, code):
        super().__init__(code)
        self.code = code


class TextBudget:
    """Stops early instead of building text far beyond the limit."""

    def __init__(self):
        self.used = 0

    def add(self, text):
        self.used += len(text) + 1
        if self.used > MAX_CHARACTERS:
            raise Unusable("TEXT_TOO_LARGE")
        return text


def flat(value):
    """One cell or run of text on a single line."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return re.sub(r"[\r\n\t]+", " ", str(value)).strip()


def extract_docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(path)
    budget = TextBudget()
    lines = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            lines.append(budget.add(Paragraph(child, document).text))
        elif tag == "tbl":
            for row in Table(child, document).rows:
                lines.append(budget.add("\t".join(flat(cell.text) for cell in row.cells)))
    return {"detected_type": "docx", "text": "\n".join(lines), "row_count": len(lines)}


def extract_xlsx(path):
    import openpyxl

    # Cached values first; a formula without a cached value (files written by openpyxl) shows the formula.
    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = values.worksheets
        if len(sheets) > MAX_SHEETS:
            raise Unusable("XLSX_SHEET_LIMIT_EXCEEDED")
        budget = TextBudget()
        sections = []
        rows_total = 0
        cells_total = 0
        rows_seen = 0
        for sheet, formula_sheet in zip(sheets, formulas.worksheets):
            lines = [budget.add(f"## {flat(sheet.title)}")]
            pairs = itertools.zip_longest(sheet.iter_rows(values_only=True), formula_sheet.iter_rows(values_only=True), fillvalue=())
            for row, formula_row in pairs:
                rows_seen += 1
                if rows_seen > MAX_SPREADSHEET_ROWS:
                    raise Unusable("XLSX_ROW_LIMIT_EXCEEDED")
                cells = [flat(value if value is not None else formula)
                         for value, formula in itertools.zip_longest(row or (), formula_row or ())]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells:
                    continue
                cells_total += sum(1 for cell in cells if cell)
                if cells_total > MAX_CELLS:
                    raise Unusable("XLSX_CELL_LIMIT_EXCEEDED")
                rows_total += 1
                lines.append(budget.add("\t".join(cells)))
            sections.append("\n".join(lines))
    finally:
        values.close()
        formulas.close()
    header = f"# XLSX 文件摘要\n工作表：{len(sections)}；数据行：{rows_total}；单元格：{cells_total}"
    return {
        "detected_type": "xlsx",
        "text": "\n\n".join([header, *sections]),
        "sheet_count": len(sections),
        "row_count": rows_total,
    }


def shape_lines(shapes, budget):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from shape_lines(shape.shapes, budget)
            continue
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = flat("".join(run.text for run in paragraph.runs))
                if text:
                    yield budget.add(text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                yield budget.add("\t".join(flat(cell.text) for cell in row.cells))


def extract_pptx(path):
    import pptx

    deck = pptx.Presentation(path)
    slides = list(deck.slides)
    if len(slides) > MAX_SLIDES:
        raise Unusable("PPTX_SLIDE_LIMIT_EXCEEDED")
    budget = TextBudget()
    sections = []
    for number, slide in enumerate(slides, start=1):
        lines = [budget.add(f"## 第 {number} 页"), *shape_lines(slide.shapes, budget)]
        if slide.has_notes_slide:
            notes = flat(slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else "")
            if notes:
                lines.append(budget.add(f"备注：{notes}"))
        sections.append("\n".join(lines))
    header = f"# PPTX 文件摘要\n幻灯片：{len(slides)}"
    return {"detected_type": "pptx", "text": "\n\n".join([header, *sections]), "page_count": len(slides)}


def extract_pdf(path):
    info = subprocess.run(["pdfinfo", path], capture_output=True, timeout=TOOL_TIMEOUT_SECONDS)
    if info.returncode != 0:
        raise Unusable("PDF_UNREADABLE")
    pages = re.search(rb"^Pages:\s+(\d+)", info.stdout, re.MULTILINE)
    page_count = int(pages.group(1)) if pages else None
    if page_count is not None and page_count > MAX_PDF_PAGES:
        raise Unusable("PDF_PAGE_LIMIT_EXCEEDED")
    result = subprocess.run(["pdftotext", "-enc", "UTF-8", path, "-"], capture_output=True, timeout=TOOL_TIMEOUT_SECONDS)
    if result.returncode != 0:
        raise Unusable("PDF_UNREADABLE")
    text = result.stdout.decode("utf-8", errors="replace").replace("\f", "\n\n")
    if len(text) > MAX_CHARACTERS * 2:
        raise Unusable("TEXT_TOO_LARGE")
    return {"detected_type": "pdf", "text": text, "page_count": page_count}


EXTRACTORS = {".docx": extract_docx, ".xlsx": extract_xlsx, ".pptx": extract_pptx, ".pdf": extract_pdf}


def write(path, content):
    with open(path, "w", encoding="utf-8") as file:
        file.write(content)


def main(names):
    for index, name in enumerate(names):
        try:
            extractor = EXTRACTORS.get(os.path.splitext(name)[1].lower())
            if extractor is None:
                raise Unusable("FILE_TYPE_UNSUPPORTED")
            result = extractor(os.path.join("/input", os.path.basename(name)))
            text = result.pop("text")
            if not text.strip():
                raise Unusable("NO_EXTRACTABLE_TEXT")
            write(f"/output/{index}.txt", text)
            meta = {"ok": True, **result}
        except Unusable as unusable:
            meta = {"ok": False, "code": unusable.code}
        except subprocess.TimeoutExpired:
            meta = {"ok": False, "code": "PARSE_TIMEOUT"}
        except Exception:  # noqa: BLE001 - any parser failure only means "no text" for this file
            meta = {"ok": False, "code": "PARSE_FAILED"}
        write(f"/output/{index}.json", json.dumps(meta))


if __name__ == "__main__":
    main(sys.argv[1:])
